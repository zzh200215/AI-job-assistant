# -*- coding: utf-8 -*-
"""Export candidate screening reports to DOCX/PDF."""

from __future__ import annotations

import html
import os
import uuid
from datetime import datetime
from typing import Any

from app.core.config import settings
from app.models.candidate_screening import CandidateScreeningSession


def _safe_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return []


def _safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return default


def _safe_score(value: Any) -> int:
    if isinstance(value, dict):
        value = value.get("score", 0)
    return _safe_int(value, 0)


def _score_level(score: int) -> str:
    if score >= 85:
        return "强匹配"
    if score >= 70:
        return "较匹配"
    if score >= 55:
        return "可关注"
    return "需谨慎"


def _join_text(items: list[str], fallback: str) -> str:
    return "、".join(items) if items else fallback


def _recommendation_distribution_text(summary: dict[str, Any]) -> str:
    distribution = summary.get("recommendation_distribution", {})
    if not isinstance(distribution, dict) or not distribution:
        return "暂无"
    ordered = sorted(
        ((str(key), _safe_int(value, 0)) for key, value in distribution.items()),
        key=lambda item: (-item[1], item[0]),
    )
    return " / ".join(f"{label} {count}人" for label, count in ordered if label)


def _pick_top_candidate(candidates: list[dict[str, Any]]) -> dict[str, Any]:
    if candidates:
        return candidates[0]
    return {}


def _dimension_rows(candidate: dict[str, Any]) -> list[dict[str, Any]]:
    dimension_scores = candidate.get("dimension_scores", {})
    if not isinstance(dimension_scores, dict):
        dimension_scores = {}

    dimensions = [
        ("skills", "技能匹配"),
        ("project", "项目经历"),
        ("experience", "工作经验"),
        ("education", "学历要求"),
        ("keyword", "关键词覆盖"),
        ("bonus", "加分项"),
    ]
    rows = []
    for key, label in dimensions:
        value = dimension_scores.get(key, {})
        if not isinstance(value, dict):
            value = {"score": value}
        rows.append(
            {
                "label": label,
                "score": _safe_score(value),
                "reason": str(value.get("reason") or value.get("summary") or "").strip(),
            }
        )
    return rows


def _candidate_context(candidate: dict[str, Any], rank: int) -> dict[str, Any]:
    matched_skills = _safe_list(candidate.get("matched_skills"))
    missing_skills = _safe_list(candidate.get("missing_required_skills"))
    risk_points = _safe_list(candidate.get("risk_points"))
    suggestions = _safe_list(candidate.get("optimization_suggestions"))
    skills = _safe_list(candidate.get("skills"))
    overall_score = _safe_score(candidate.get("overall_score"))

    return {
        "rank": rank,
        "resume_id": candidate.get("resume_id"),
        "candidate_name": str(candidate.get("candidate_name") or f"候选人#{rank}"),
        "file_name": str(candidate.get("file_name") or ""),
        "years_exp": _safe_int(candidate.get("years_exp"), 0),
        "overall_score": overall_score,
        "score_level": _score_level(overall_score),
        "recommendation": str(candidate.get("recommendation") or "待评估"),
        "overall_reason": str(candidate.get("overall_reason") or "").strip(),
        "skills": skills,
        "skills_text": _join_text(skills, "未提取"),
        "matched_skills": matched_skills,
        "matched_skills_text": _join_text(matched_skills, "暂无明显命中"),
        "missing_skills": missing_skills,
        "missing_skills_text": _join_text(missing_skills, "无明显硬缺口"),
        "risk_points": risk_points,
        "suggestions": suggestions,
        "dimension_rows": _dimension_rows(candidate),
    }


def _build_report_context(session: CandidateScreeningSession) -> dict[str, Any]:
    result = session.result_payload or {}
    summary = result.get("summary", {}) if isinstance(result, dict) else {}
    raw_candidates = result.get("candidates", []) if isinstance(result, dict) else []
    candidates = [
        _candidate_context(candidate, rank=index)
        for index, candidate in enumerate(raw_candidates, start=1)
        if isinstance(candidate, dict)
    ]
    top_candidate = _pick_top_candidate(candidates)
    skill_gaps = summary.get("most_common_skill_gaps", [])
    if not isinstance(skill_gaps, list):
        skill_gaps = []

    report_title = session.name or "候选人筛选对比报告"
    jd_title = str(summary.get("jd_title") or session.jd_title or "未命名岗位")
    company = str(summary.get("company") or session.company or "未填写")
    candidate_count = _safe_int(summary.get("returned_candidates"), len(candidates))
    total_candidates = _safe_int(summary.get("total_candidates"), session.candidate_count or len(candidates))
    average_score = round(
        sum(candidate["overall_score"] for candidate in candidates) / len(candidates),
        1,
    ) if candidates else 0

    gap_rows = []
    for item in skill_gaps[:8]:
        if not isinstance(item, dict):
            continue
        skill = str(item.get("skill") or "").strip()
        if not skill:
            continue
        count = _safe_int(item.get("count"), 0)
        ratio = f"{round((count / max(total_candidates, 1)) * 100)}%"
        gap_rows.append({"skill": skill, "count": count, "ratio": ratio})

    return {
        "report_title": report_title,
        "jd_title": jd_title,
        "company": company,
        "candidate_count": candidate_count,
        "total_candidates": total_candidates,
        "average_score": average_score,
        "recommendation_distribution_text": _recommendation_distribution_text(summary),
        "top_candidate": top_candidate,
        "skill_gaps": gap_rows,
        "candidates": candidates,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def _create_export_path(session_id: int, ext: str) -> tuple[str, str]:
    stored_name = f"screening_{session_id}_{uuid.uuid4().hex}.{ext}"
    now = datetime.now()
    rel_dir = os.path.join(settings.UPLOAD_DIR, "export", str(now.year), f"{now.month:02d}")
    abs_dir = os.path.abspath(rel_dir)
    os.makedirs(abs_dir, exist_ok=True)
    abs_path = os.path.join(abs_dir, stored_name)
    rel_path = os.path.join(rel_dir, stored_name).replace("\\", "/")
    return abs_path, rel_path


def _set_cell_text(cell, text: str, *, bold: bool = False, font_size: int = 10):
    from docx.shared import Pt

    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)


def _add_summary_table(doc, context: dict[str, Any]):
    from docx.shared import Pt, RGBColor

    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    cells = table.rows[0].cells
    _set_cell_text(cells[0], "岗位", bold=True)
    _set_cell_text(cells[1], "公司", bold=True)
    _set_cell_text(cells[2], "导出时间", bold=True)

    values = table.rows[1].cells
    _set_cell_text(values[0], context["jd_title"])
    _set_cell_text(values[1], context["company"])
    _set_cell_text(values[2], context["generated_at"])

    table2 = doc.add_table(rows=2, cols=4)
    table2.style = "Table Grid"
    headers = table2.rows[0].cells
    for cell, label in zip(headers, ["纳入候选人", "实际返回", "平均分", "推荐分布"]):
        _set_cell_text(cell, label, bold=True)

    values2 = table2.rows[1].cells
    _set_cell_text(values2[0], f"{context['total_candidates']} 人")
    _set_cell_text(values2[1], f"{context['candidate_count']} 人")
    _set_cell_text(values2[2], f"{context['average_score']} 分")
    _set_cell_text(values2[3], context["recommendation_distribution_text"])

    p = doc.add_paragraph()
    run = p.add_run("报告摘要")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 102, 204)


def _add_top_candidate_block(doc, context: dict[str, Any]):
    from docx.shared import Pt

    top_candidate = context.get("top_candidate") or {}
    doc.add_heading("Top 候选人概览", level=2)
    if not top_candidate:
        doc.add_paragraph("暂无可展示候选人。")
        return

    paragraph = doc.add_paragraph()
    lead = paragraph.add_run(
        f"#{top_candidate['rank']} {top_candidate['candidate_name']} | {top_candidate['overall_score']} 分 | {top_candidate['recommendation']}"
    )
    lead.bold = True
    lead.font.size = Pt(12)
    doc.add_paragraph(f"工作年限：{top_candidate['years_exp']} 年")
    doc.add_paragraph(f"命中技能：{top_candidate['matched_skills_text']}")
    doc.add_paragraph(f"核心缺口：{top_candidate['missing_skills_text']}")
    if top_candidate["overall_reason"]:
        doc.add_paragraph(f"综合判断：{top_candidate['overall_reason']}")


def _add_skill_gap_table(doc, context: dict[str, Any]):
    doc.add_heading("共性技能缺口", level=2)
    rows = context.get("skill_gaps") or []
    if not rows:
        doc.add_paragraph("本次筛选未发现显著共性缺口。")
        return

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["技能项", "出现次数", "占比"]):
        _set_cell_text(cell, label, bold=True)

    for row in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], row["skill"])
        _set_cell_text(cells[1], f"{row['count']} 次")
        _set_cell_text(cells[2], row["ratio"])


def _add_candidate_section(doc, candidate: dict[str, Any]):
    from docx.shared import Pt

    doc.add_heading(
        f"#{candidate['rank']} {candidate['candidate_name']} | {candidate['overall_score']} 分 | {candidate['score_level']}",
        level=2,
    )

    meta = doc.add_paragraph()
    meta.add_run("推荐结论：").bold = True
    meta.add_run(candidate["recommendation"])
    meta.add_run("    ")
    meta.add_run("工作年限：").bold = True
    meta.add_run(f"{candidate['years_exp']} 年")
    if candidate["file_name"]:
        meta.add_run("    ")
        meta.add_run("文件：").bold = True
        meta.add_run(candidate["file_name"])

    summary = doc.add_paragraph()
    summary.add_run("技能画像：").bold = True
    summary.add_run(candidate["skills_text"])

    hits = doc.add_paragraph()
    hits.add_run("命中技能：").bold = True
    hits.add_run(candidate["matched_skills_text"])

    gaps = doc.add_paragraph()
    gaps.add_run("缺失技能：").bold = True
    gaps.add_run(candidate["missing_skills_text"])

    if candidate["overall_reason"]:
        reason = doc.add_paragraph()
        reason.add_run("综合判断：").bold = True
        run = reason.add_run(candidate["overall_reason"])
        run.font.size = Pt(10.5)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["评估维度", "得分", "说明"]):
        _set_cell_text(cell, label, bold=True)

    for row in candidate["dimension_rows"]:
        cells = table.add_row().cells
        _set_cell_text(cells[0], row["label"])
        _set_cell_text(cells[1], str(row["score"]))
        _set_cell_text(cells[2], row["reason"] or "-")

    if candidate["risk_points"]:
        doc.add_paragraph("风险提示", style=None).runs[0].bold = True
        for item in candidate["risk_points"]:
            doc.add_paragraph(item, style="List Bullet")

    if candidate["suggestions"]:
        doc.add_paragraph("跟进建议", style=None).runs[0].bold = True
        for item in candidate["suggestions"]:
            doc.add_paragraph(item, style="List Bullet")


def export_screening_docx(session: CandidateScreeningSession) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    context = _build_report_context(session)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(context["report_title"])
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(34, 62, 105)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f"{context['jd_title']} | {context['company']}")
    subtitle_run.font.size = Pt(11)

    doc.add_paragraph("")
    _add_summary_table(doc, context)
    _add_top_candidate_block(doc, context)
    _add_skill_gap_table(doc, context)

    doc.add_heading("候选人明细", level=1)
    if not context["candidates"]:
        doc.add_paragraph("暂无候选人数据。")
    else:
        for candidate in context["candidates"]:
            _add_candidate_section(doc, candidate)

    abs_path, rel_path = _create_export_path(session.id, "docx")
    doc.save(abs_path)
    return rel_path


def _candidate_card_html(candidate: dict[str, Any]) -> str:
    dimension_rows = "".join(
        f"""
        <tr>
          <td>{html.escape(row['label'])}</td>
          <td>{row['score']}</td>
          <td>{html.escape(row['reason'] or '-')}</td>
        </tr>
        """
        for row in candidate["dimension_rows"]
    )

    risk_html = ""
    if candidate["risk_points"]:
        items = "".join(f"<li>{html.escape(item)}</li>" for item in candidate["risk_points"])
        risk_html = f"<div class='sub-block'><div class='sub-title'>风险提示</div><ul>{items}</ul></div>"

    suggestion_html = ""
    if candidate["suggestions"]:
        items = "".join(f"<li>{html.escape(item)}</li>" for item in candidate["suggestions"])
        suggestion_html = f"<div class='sub-block'><div class='sub-title'>跟进建议</div><ul>{items}</ul></div>"

    overall_reason = ""
    if candidate["overall_reason"]:
        overall_reason = (
            "<div class='sub-block'><div class='sub-title'>综合判断</div>"
            f"<p>{html.escape(candidate['overall_reason'])}</p></div>"
        )

    return f"""
    <section class="candidate-card">
      <div class="candidate-head">
        <div>
          <h3>#{candidate['rank']} {html.escape(candidate['candidate_name'])}</h3>
          <p class="meta">{html.escape(candidate['recommendation'])} · {candidate['years_exp']} 年经验 · {html.escape(candidate['file_name'] or '未记录文件')}</p>
        </div>
        <div class="score-box">
          <span class="score">{candidate['overall_score']}</span>
          <span class="score-level">{html.escape(candidate['score_level'])}</span>
        </div>
      </div>
      <div class="chips">
        <span>命中技能：{html.escape(candidate['matched_skills_text'])}</span>
        <span>缺失技能：{html.escape(candidate['missing_skills_text'])}</span>
      </div>
      <div class="sub-block">
        <div class="sub-title">技能画像</div>
        <p>{html.escape(candidate['skills_text'])}</p>
      </div>
      {overall_reason}
      <table class="score-table">
        <thead>
          <tr><th>评估维度</th><th>得分</th><th>说明</th></tr>
        </thead>
        <tbody>
          {dimension_rows}
        </tbody>
      </table>
      {risk_html}
      {suggestion_html}
    </section>
    """


def _render_screening_html(context: dict[str, Any]) -> str:
    top_candidate = context.get("top_candidate") or {}
    top_candidate_html = "<p>暂无可展示候选人。</p>"
    if top_candidate:
        top_candidate_html = f"""
        <div class="hero-top-card">
          <div>
            <div class="hero-label">Top 候选人</div>
            <div class="hero-name">#{top_candidate['rank']} {html.escape(top_candidate['candidate_name'])}</div>
            <p>{html.escape(top_candidate['recommendation'])} · {top_candidate['years_exp']} 年经验</p>
          </div>
          <div class="hero-score">
            <span>{top_candidate['overall_score']}</span>
            <small>{html.escape(top_candidate['score_level'])}</small>
          </div>
        </div>
        <div class="hero-details">
          <p><strong>命中技能：</strong>{html.escape(top_candidate['matched_skills_text'])}</p>
          <p><strong>缺失技能：</strong>{html.escape(top_candidate['missing_skills_text'])}</p>
          <p><strong>综合判断：</strong>{html.escape(top_candidate['overall_reason'] or '详见下方候选人明细')}</p>
        </div>
        """

    gap_rows = context.get("skill_gaps") or []
    if gap_rows:
        gap_html = "".join(
            f"<tr><td>{html.escape(row['skill'])}</td><td>{row['count']} 次</td><td>{html.escape(row['ratio'])}</td></tr>"
            for row in gap_rows
        )
    else:
        gap_html = "<tr><td colspan='3'>本次筛选未发现显著共性缺口。</td></tr>"

    candidates_html = "".join(_candidate_card_html(candidate) for candidate in context["candidates"])
    if not candidates_html:
        candidates_html = "<p>暂无候选人数据。</p>"

    return f"""
    <!DOCTYPE html>
    <html lang="zh-CN">
    <head>
      <meta charset="utf-8">
      <style>
        @page {{
          size: A4;
          margin: 16mm 14mm;
        }}
        body {{
          font-family: "Microsoft YaHei", "PingFang SC", sans-serif;
          color: #213047;
          line-height: 1.55;
          font-size: 12px;
        }}
        h1, h2, h3, p {{
          margin: 0;
        }}
        .hero {{
          background: linear-gradient(135deg, #f4f7fb, #e8f0ff);
          border: 1px solid #d7e3f8;
          border-radius: 14px;
          padding: 20px 22px;
          margin-bottom: 18px;
        }}
        .hero h1 {{
          font-size: 24px;
          color: #203864;
          margin-bottom: 6px;
        }}
        .hero-sub {{
          color: #5f6f86;
          margin-bottom: 16px;
        }}
        .summary-grid {{
          display: grid;
          grid-template-columns: repeat(4, 1fr);
          gap: 10px;
          margin-bottom: 16px;
        }}
        .summary-item {{
          background: #fff;
          border-radius: 10px;
          padding: 12px;
          border: 1px solid #dbe5f3;
        }}
        .summary-item .label {{
          display: block;
          font-size: 10px;
          color: #70819a;
          margin-bottom: 4px;
        }}
        .summary-item .value {{
          font-size: 16px;
          font-weight: 700;
          color: #203864;
        }}
        .hero-top-card {{
          display: flex;
          justify-content: space-between;
          gap: 12px;
          align-items: center;
          background: #fff;
          border-radius: 12px;
          border: 1px solid #dbe5f3;
          padding: 14px 16px;
          margin-bottom: 10px;
        }}
        .hero-label {{
          font-size: 10px;
          color: #70819a;
          margin-bottom: 4px;
        }}
        .hero-name {{
          font-size: 18px;
          font-weight: 700;
          color: #203864;
        }}
        .hero-score {{
          min-width: 88px;
          text-align: center;
          background: #203864;
          color: #fff;
          border-radius: 12px;
          padding: 12px 10px;
        }}
        .hero-score span {{
          display: block;
          font-size: 28px;
          font-weight: 700;
          line-height: 1;
        }}
        .hero-score small {{
          display: block;
          margin-top: 5px;
          font-size: 11px;
        }}
        .hero-details p {{
          margin-top: 4px;
        }}
        .section {{
          margin-bottom: 18px;
        }}
        .section h2 {{
          color: #203864;
          font-size: 16px;
          margin-bottom: 10px;
          padding-bottom: 5px;
          border-bottom: 1px solid #dbe5f3;
        }}
        table {{
          width: 100%;
          border-collapse: collapse;
        }}
        .gap-table th, .gap-table td, .score-table th, .score-table td {{
          border: 1px solid #dbe5f3;
          padding: 7px 8px;
          vertical-align: top;
        }}
        .gap-table th, .score-table th {{
          background: #f4f7fb;
          color: #203864;
          text-align: left;
        }}
        .candidate-card {{
          border: 1px solid #dbe5f3;
          border-radius: 12px;
          padding: 16px;
          margin-bottom: 14px;
          page-break-inside: avoid;
        }}
        .candidate-head {{
          display: flex;
          justify-content: space-between;
          gap: 12px;
          align-items: center;
          margin-bottom: 10px;
        }}
        .candidate-head h3 {{
          font-size: 16px;
          color: #203864;
          margin-bottom: 4px;
        }}
        .meta {{
          color: #6d7d92;
          font-size: 11px;
        }}
        .score-box {{
          min-width: 74px;
          text-align: center;
          border-radius: 10px;
          background: #eff4ff;
          padding: 10px 8px;
          color: #203864;
        }}
        .score {{
          display: block;
          font-size: 24px;
          font-weight: 700;
          line-height: 1;
        }}
        .score-level {{
          display: block;
          margin-top: 4px;
          font-size: 11px;
        }}
        .chips {{
          margin-bottom: 10px;
        }}
        .chips span {{
          display: inline-block;
          background: #f4f7fb;
          border-radius: 999px;
          padding: 5px 10px;
          margin: 0 8px 8px 0;
          color: #42546b;
        }}
        .sub-block {{
          margin-bottom: 10px;
        }}
        .sub-title {{
          font-weight: 700;
          color: #203864;
          margin-bottom: 4px;
        }}
        ul {{
          margin: 6px 0 0 18px;
          padding: 0;
        }}
        li {{
          margin-bottom: 4px;
        }}
      </style>
    </head>
    <body>
      <section class="hero">
        <h1>{html.escape(context['report_title'])}</h1>
        <p class="hero-sub">{html.escape(context['jd_title'])} | {html.escape(context['company'])} | 导出时间 {html.escape(context['generated_at'])}</p>
        <div class="summary-grid">
          <div class="summary-item"><span class="label">纳入候选人</span><span class="value">{context['total_candidates']} 人</span></div>
          <div class="summary-item"><span class="label">实际返回</span><span class="value">{context['candidate_count']} 人</span></div>
          <div class="summary-item"><span class="label">平均分</span><span class="value">{context['average_score']} 分</span></div>
          <div class="summary-item"><span class="label">推荐分布</span><span class="value" style="font-size:13px">{html.escape(context['recommendation_distribution_text'])}</span></div>
        </div>
        {top_candidate_html}
      </section>

      <section class="section">
        <h2>共性技能缺口</h2>
        <table class="gap-table">
          <thead><tr><th>技能项</th><th>出现次数</th><th>占比</th></tr></thead>
          <tbody>{gap_html}</tbody>
        </table>
      </section>

      <section class="section">
        <h2>候选人明细</h2>
        {candidates_html}
      </section>
    </body>
    </html>
    """


def export_screening_pdf(session: CandidateScreeningSession) -> str:
    try:
        from weasyprint import HTML as WeasyprintHTML
    except ImportError:
        raise RuntimeError("PDF导出需要安装 weasyprint: pip install weasyprint")

    context = _build_report_context(session)
    html_content = _render_screening_html(context)
    abs_path, rel_path = _create_export_path(session.id, "pdf")
    WeasyprintHTML(string=html_content).write_pdf(abs_path)
    return rel_path
