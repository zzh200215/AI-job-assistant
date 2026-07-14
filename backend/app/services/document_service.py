"""文档解析服务：支持 txt / md / pdf / docx 转纯文本"""

import os


def _parse_txt(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def _parse_md(path: str) -> str:
    """Markdown: 保留文本结构，移除图片链接等"""
    import re

    with open(path, encoding="utf-8", errors="ignore") as f:
        text = f.read()
    # 移除图片 ![](url)
    text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
    # 移除行内链接 [text](url) 但保留 text
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def _parse_pdf(path: str) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t:
                text_parts.append(t)
    return "\n".join(text_parts)


def _parse_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_document(path: str, file_type: str | None = None) -> str:
    """
    根据文件类型解析为纯文本。
    支持: txt, md, pdf, docx
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f"文件不存在: {path}")

    ft = (file_type or os.path.splitext(path)[1].lstrip(".")).lower()

    if ft == "txt":
        return _parse_txt(path)
    elif ft == "md":
        return _parse_md(path)
    elif ft == "pdf":
        return _parse_pdf(path)
    elif ft in ("docx", "doc"):
        return _parse_docx(path)
    else:
        raise ValueError(f"不支持的文件类型: {ft}")
