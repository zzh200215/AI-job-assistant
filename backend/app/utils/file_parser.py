"""Resume file parsing utilities with optional OCR support."""

from __future__ import annotations

import os
import re
import shutil
from importlib import import_module


def _load_optional_module(module_name: str):
    try:
        return import_module(module_name)
    except Exception:
        return None


def _configure_tesseract(pytesseract_module) -> bool:
    configured_cmd = os.getenv("TESSERACT_CMD", "").strip()
    if configured_cmd:
        pytesseract_module.pytesseract.tesseract_cmd = configured_cmd
        return os.path.exists(configured_cmd)

    binary_name = getattr(pytesseract_module.pytesseract, "tesseract_cmd", "tesseract") or "tesseract"
    return shutil.which(binary_name) is not None


def is_ocr_available() -> bool:
    """Return whether the current runtime can OCR scanned PDFs."""
    pytesseract_module = _load_optional_module("pytesseract")
    if pytesseract_module is None:
        return False
    if not _configure_tesseract(pytesseract_module):
        return False

    image_backend_ready = (
        _load_optional_module("pypdfium2") is not None or _load_optional_module("pdf2image") is not None
    )
    return image_backend_ready


def _render_pdf_images(path: str):
    pdfium_module = _load_optional_module("pypdfium2")
    if pdfium_module is not None:
        pdf = pdfium_module.PdfDocument(path)
        try:
            for page_index in range(len(pdf)):
                page = pdf[page_index]
                try:
                    yield page.render(scale=2).to_pil()
                finally:
                    page.close()
        finally:
            pdf.close()
        return

    pdf2image_module = _load_optional_module("pdf2image")
    if pdf2image_module is None:
        raise ValueError("当前环境未安装 PDF OCR 渲染依赖")

    yield from pdf2image_module.convert_from_path(path, dpi=200)


def _ocr_pdf(path: str) -> str:
    if not is_ocr_available():
        raise ValueError("当前环境未开启 OCR 识别能力")

    pytesseract_module = import_module("pytesseract")
    _configure_tesseract(pytesseract_module)

    texts = []
    try:
        for image in _render_pdf_images(path):
            try:
                text = pytesseract_module.image_to_string(
                    image,
                    lang=os.getenv("OCR_LANG", "chi_sim+eng"),
                )
            finally:
                close_image = getattr(image, "close", None)
                if callable(close_image):
                    close_image()
            if text and text.strip():
                texts.append(text.strip())
    except pytesseract_module.TesseractNotFoundError as exc:
        raise ValueError("未找到 Tesseract OCR 可执行文件") from exc
    except Exception as exc:
        raise ValueError(f"OCR 识别失败: {exc}") from exc

    full_text = "\n".join(texts)
    if not full_text.strip():
        raise ValueError("扫描件 OCR 未识别出可用文本")
    return full_text


def _ocr_image(path: str) -> str:
    if not is_ocr_available():
        raise ValueError("当前环境未开启 OCR 识别能力")

    pytesseract_module = import_module("pytesseract")
    _configure_tesseract(pytesseract_module)
    image_module = _load_optional_module("PIL.Image")
    if image_module is None:
        raise ValueError("当前环境未安装图片 OCR 依赖")

    try:
        image = image_module.open(path)
        try:
            text = pytesseract_module.image_to_string(
                image,
                lang=os.getenv("OCR_LANG", "chi_sim+eng"),
            )
        finally:
            close_image = getattr(image, "close", None)
            if callable(close_image):
                close_image()
    except pytesseract_module.TesseractNotFoundError as exc:
        raise ValueError("未找到 Tesseract OCR 可执行文件") from exc
    except Exception as exc:
        raise ValueError(f"图片 OCR 识别失败: {exc}") from exc

    if not text or not text.strip():
        raise ValueError("图片 OCR 未识别出可用文本")
    return text.strip()


def _parse_pdf(path: str) -> str:
    """Parse PDF and fall back to OCR for scanned PDFs when available."""
    import pdfplumber

    try:
        with pdfplumber.open(path) as pdf:
            if pdf.metadata and pdf.metadata.get("encrypted"):
                raise ValueError("encrypted")

            text_parts = []
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text.strip():
                    text_parts.append(text.strip())
    except ValueError as exc:
        if "encrypted" in str(exc):
            raise ValueError("该 PDF 文件已加密，请先解密后重新上传") from exc
        raise
    except Exception as exc:
        err = str(exc).lower()
        if "encrypted" in err or "password" in err:
            raise ValueError("该 PDF 文件已加密，请先解密后重新上传") from exc
        if "not a pdf" in err or "corrupt" in err or "cannot read" in err:
            raise ValueError("PDF 文件格式异常或已损坏，请检查后重新上传") from exc
        raise ValueError(f"PDF 解析失败: {exc}") from exc

    full_text = "\n".join(text_parts)
    if full_text.strip():
        return full_text

    if is_ocr_available():
        return _ocr_pdf(path)
    raise ValueError("该 PDF 文件为扫描件或图片型 PDF，当前环境未开启 OCR，请上传可编辑 PDF 或 DOCX")


def _parse_docx(path: str) -> str:
    """Parse DOCX paragraphs and tables."""
    from docx import Document

    try:
        doc = Document(path)
    except Exception as exc:
        raise ValueError(f"DOCX 文件解析失败: {exc}") from exc

    parts = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        if table_index > 1 or parts:
            parts.append("")
        parts.append(f"[表格 {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
        parts.append("")

    full_text = "\n".join(parts)
    if not full_text.strip():
        raise ValueError("DOCX 文件内容为空，请检查文件")
    return full_text


def _parse_txt(path: str) -> str:
    try:
        with open(path, encoding="utf-8", errors="ignore") as file:
            text = file.read()
    except Exception as exc:
        raise ValueError(f"TXT 文件解析失败: {exc}") from exc

    if not text.strip():
        raise ValueError("TXT 文件内容为空")
    return text


def _parse_md(path: str) -> str:
    try:
        with open(path, encoding="utf-8", errors="ignore") as file:
            text = file.read()
    except Exception as exc:
        raise ValueError(f"Markdown 文件解析失败: {exc}") from exc

    text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    if not text.strip():
        raise ValueError("Markdown 文件内容为空")
    return text


def parse_resume(path: str, file_type: str | None = None) -> str:
    """Parse supported resume file types into plain text."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"文件不存在: {path}")

    try:
        if os.path.getsize(path) == 0:
            raise ValueError("文件内容为空")
    except OSError as exc:
        raise ValueError(f"无法读取文件: {exc}") from exc

    file_ext = (file_type or os.path.splitext(path)[1].lstrip(".")).lower()
    parsers = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "doc": _parse_docx,
        "txt": _parse_txt,
        "md": _parse_md,
        "png": _ocr_image,
        "jpg": _ocr_image,
        "jpeg": _ocr_image,
    }

    parser = parsers.get(file_ext)
    if parser is None:
        raise ValueError(f"不支持的文件类型 .{file_ext}")

    if file_ext == "doc":
        raise ValueError("检测到 .doc 旧版 Word 文件，请另存为 .docx 后重新上传")

    return parser(path)
