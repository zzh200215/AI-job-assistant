# -*- coding: utf-8 -*-
"""Regenerate the target resume docx with a tighter LLM application engineer focus."""
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "曾子豪_大模型应用开发工程师.docx"
BACKUP = ROOT / "曾子豪_大模型应用开发工程师.原始备份.docx"


def configure_document(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "微软雅黑"
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.22

    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)


def add_section_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12.8)
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    p_pr = p._p.get_or_add_pPr()
    border = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        '<w:bottom w:val="single" w:sz="8" w:space="1" w:color="1A3A5C"/>'
        '</w:pBdr>'
    )
    p_pr.append(border)


def add_paragraph(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    size: float = 10.5,
    color: RGBColor | None = None,
    align=None,
) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if align is not None:
        p.alignment = align


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.55)


def add_project_header(doc: Document, name: str, time: str, tech: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"◎ {name}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    time_run = p.add_run(f"    {time}")
    time_run.font.size = Pt(9)
    time_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    run2 = p2.add_run(f"技术栈：{tech}")
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def build_resume() -> Document:
    doc = Document()
    configure_document(doc)

    add_paragraph(
        doc,
        "曾子豪",
        bold=True,
        size=22,
        color=RGBColor(0x1A, 0x3A, 0x5C),
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(
        doc,
        "求职意向：大模型应用开发工程师 / AI 应用开发工程师",
        size=10,
        color=RGBColor(0x66, 0x66, 0x66),
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(
        doc,
        "17779699791  |  2974467965@qq.com  |  江西吉安  |  2025届本科",
        size=9.2,
        color=RGBColor(0x44, 0x44, 0x44),
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_section_title(doc, "核心优势")
    for item in [
        "围绕 RAG、Agent 编排、Prompt Engineering、LLM API 集成做过完整项目落地，能把模型能力接成可用系统。",
        "具备 Python 后端全栈实现能力，熟悉 FastAPI、WebSocket、SQLAlchemy、MySQL、Chroma 等工程组件。",
        "项目表达偏工程落地而非概念堆砌，能独立完成需求拆解、链路设计、异常兜底和结果交付。",
    ]:
        add_bullet(doc, item)

    add_section_title(doc, "教育背景")
    p = doc.add_paragraph()
    run = p.add_run("南昌大学科学技术学院")
    run.bold = True
    run.font.size = Pt(10.5)
    run2 = p.add_run("    计算机科学与技术 本科    2021.09 - 2025.06")
    run2.font.size = Pt(10.5)
    run2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    add_paragraph(
        doc,
        "核心课程：数据结构与算法、操作系统、计算机网络、数据库原理、Python 程序设计、机器学习基础",
        size=9.5,
        color=RGBColor(0x44, 0x44, 0x44),
    )

    add_section_title(doc, "专业技能")
    skills = [
        (
            "大模型应用：",
            "RAG、Prompt Engineering、Agent 编排、LangChain、LangGraph、DeepAgents、Query Rewrite、Embedding、OpenAI / Qwen API 集成",
        ),
        (
            "后端与工程：",
            "Python、FastAPI、RESTful API、WebSocket、SQLAlchemy、MySQL、Redis、Docker、Git",
        ),
        (
            "检索与数据：",
            "Chroma 向量检索、多路召回、规则打分、结构化数据查询、知识库管理",
        ),
        (
            "基础能力：",
            "C / C++、Linux 系统编程、Socket / TCP/IP、Shell、JavaScript、Vue3（基础协作开发）",
        ),
    ]
    for category, content in skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(category)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
        run2 = p.add_run(content)
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    add_section_title(doc, "项目经历")

    add_project_header(
        doc,
        "智能招聘与职业规划平台（Agentic RAG 多智能体项目）",
        "2025.03 - 至今",
        "Python、FastAPI、LangChain、LangGraph、Chroma、Qwen、Vue3、WebSocket",
    )
    add_paragraph(
        doc,
        "独立完成一套面向求职场景的 AI 工作台，覆盖简历解析、岗位理解、RAG 检索、匹配分析、职业规划、模拟面试和历史复盘等主链路。",
        size=9.5,
    )
    for item in [
        "设计 Agentic RAG 分析链路，将意图识别、解析、检索、匹配、优化、面试题生成、职业规划和报告汇总串成统一流程，并加入关键步骤降级与异常兜底。",
        "实现多智能体协作编排，按依赖关系调度简历诊断、岗位分析、匹配评估、面试辅导、职业规划和汇总智能体，支持动态选择分析路径。",
        "搭建 Query Rewrite 检索改写服务，把用户原始问题改写为多维查询后并行召回 Chroma 知识片段，再做合并去重，提升检索覆盖面和结果可用性。",
        "实现规则引擎 + LLM 解释的匹配分析方式，由规则负责技能、经验、关键词等维度打分，LLM 负责输出自然语言解释，降低纯模型评分的不稳定性。",
        "基于 FastAPI WebSocket 实现模拟面试能力，支持实时问答、超时控制、过程评分和低分追问，使分析结果能继续流向面试准备环节。",
    ]:
        add_bullet(doc, item)

    add_project_header(
        doc,
        "DeepSearch 对话式多智能体研究系统",
        "2025.03 - 至今",
        "Python、DeepAgents、LangGraph、FastAPI、WebSocket、Tavily、MySQL、RAGFlow、React",
    )
    add_paragraph(
        doc,
        "独立开发面向复杂问题研究的多智能体系统，整合互联网搜索、数据库查询和私有知识库检索，并支持报告生成与文件交付。",
        size=9.5,
    )
    for item in [
        "设计一主三从的多智能体结构，由主智能体负责任务规划和调度，子智能体分别处理网络搜索、结构化数据查询和知识库检索。",
        "打通 Tavily 搜索、MySQL 查询和 RAGFlow 问答三类检索路径，支持多来源交叉验证，增强研究结果的完整性。",
        "基于 FastAPI WebSocket 实现任务执行过程的实时推送，把工具调用、子任务状态和异常事件同步给前端，提升可观测性。",
        "实现文件交付链路，支持读取上传附件、生成 Markdown 研究报告并转换为 PDF，形成从研究到输出的完整闭环。",
    ]:
        add_bullet(doc, item)

    add_project_header(
        doc,
        "基于 Qt 的即时通讯系统（集成 ChatGPT）",
        "2025.02 - 2025.04",
        "C++、Boost.Asio、Qt、gRPC、Redis、SQLite、ChatGPT API",
    )
    add_paragraph(
        doc,
        "完成一个跨平台 IM 系统，实现高并发通信、服务拆分和大模型对话能力接入，体现网络基础与工程实现能力。",
        size=9.5,
    )
    for item in [
        "使用 Boost.Asio 构建异步通信模块，支持长连接和高并发消息收发。",
        "拆分网关、鉴权、状态和聊天服务，基于 gRPC 与 Redis 完成服务间通信和状态缓存。",
        "接入 ChatGPT API 实现智能对话功能，补充了从传统系统到大模型应用接入的工程经验。",
    ]:
        add_bullet(doc, item)

    add_section_title(doc, "自我评价")
    for item in [
        "对大模型应用开发的理解偏工程落地，能把检索、编排、接口、前后端联调和交付结果连成完整产品链路。",
        "具备扎实的计算机基础和较强的自驱学习能力，能够快速吸收新框架并在项目中验证。",
        "适合大模型应用开发、AI 产品工程化、智能体平台、RAG 系统等方向的岗位。",
    ]:
        add_bullet(doc, item)

    return doc


def main() -> None:
    if TARGET.exists() and not BACKUP.exists():
        shutil.copy2(TARGET, BACKUP)

    doc = build_resume()
    doc.save(TARGET)
    print(f"optimized: {TARGET.name}")
    if BACKUP.exists():
        print(f"backup: {BACKUP.name}")


if __name__ == "__main__":
    main()
